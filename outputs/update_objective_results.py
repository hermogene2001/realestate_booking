from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DOCX_PATH = Path(r"E:\realestate_booking\bachelor_of_technology_capstone_project.docx")


def set_paragraph_text(paragraph, text):
    for run in list(paragraph.runs):
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p = new_p
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def insert_table_after(paragraph, rows, cols):
    table = paragraph._parent.add_table(rows=rows, cols=cols, width=Inches(6.5))
    paragraph._p.addnext(table._tbl)
    return table


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


doc = Document(DOCX_PATH)

# Clean and strengthen the Chapter Five results narrative.
replacements = {
    527: (
        "This chapter presents the implementation results of the decentralized real estate booking system and "
        "directly evaluates how the four research objectives stated in Chapter One were achieved. The results "
        "cover vulnerability analysis, smart-contract escrow implementation, user-facing dApp delivery, and "
        "technical evaluation evidence. Quantitatively, the implemented platform includes 4 running service "
        "layers, 15 backend route files, more than 80 API endpoints, 27 Prisma database models, 32 frontend "
        "pages, 3 supported payment methods, 6 automated fraud-detection rules, and a verified on-chain escrow "
        "submission."
    ),
    529: (
        "The implemented platform uses a full-stack architecture consisting of a Next.js/React frontend, an "
        "Express and TypeScript backend, a MySQL database managed through Prisma, and a Hardhat Ethereum "
        "environment for smart-contract escrow testing. The stack supports a production PWA build, REST APIs, "
        "database persistence, blockchain escrow verification, and administrative monitoring."
    ),
    531: "",
    533: (
        "The RealEstateEscrow smart contract was deployed to a Hardhat localhost Ethereum network and verified "
        "through an on-chain escrow submission. The Solidity escrow contract contains 393 lines of code and "
        "implements a booking state machine, dual tenant-owner handover confirmation, timeout refund, dispute "
        "resolution, and a configurable 2.5 percent platform fee."
    ),
    535: "",
    537: (
        "The backend exposes 15 route files and more than 80 endpoints covering authentication, properties, "
        "bookings, payments, notifications, KYC, wishlist, search, messaging, AI chat, reviews, and admin "
        "operations. This API surface supports both ordinary users and administrators, allowing the technical "
        "prototype to operate as a complete booking platform rather than a smart-contract-only demonstration."
    ),
    539: "",
    541: (
        "The MySQL database contains 27 Prisma models covering users, properties, bookings, payments, reviews, "
        "notifications, fraud alerts, disputes, commissions, documents, messages, favorites, verifications, and "
        "other supporting entities. The seed data includes 6 properties across 3 Kigali districts: Gasabo, "
        "Kicukiro, and Nyarugenge."
    ),
    543: (
        "The system supports 3 payment methods: direct Ethereum payment, MTN Mobile Money, and card payment. "
        "All payment flows are designed to culminate in a recorded escrow state so that the tenant's deposit is "
        "not dependent only on informal trust between tenant, owner, or broker."
    ),
    545: "",
    547: (
        "The security architecture uses multiple layers: JWT authentication, role-based access control for "
        "TENANT, OWNER, and ADMIN users, Zod input validation, 2FA support, KYC document submission, file-upload "
        "validation, rate limiting, fraud detection, and blockchain protections such as ReentrancyGuard and "
        "Pausable smart-contract controls."
    ),
    549: "",
    551: (
        "Three background jobs support operational reliability: booking timeout cancellation every 15 minutes, "
        "daily timeout-warning notifications at 09:00 Kigali time, and hourly fraud-detection scans. The fraud "
        "module applies 6 rule-based checks, including rapid bookings, rapid cancellations, quick handovers, "
        "duplicate wallets, price manipulation, and same-IP activity."
    ),
    553: "",
    555: (
        "Code-quality results show measurable improvement. Backend TypeScript compilation errors were reduced "
        "from approximately 60 errors to zero across 10 categories, and the frontend production build generated "
        "32 pages with a PWA service worker. These results show that the prototype is not only conceptually "
        "designed but also technically buildable and testable."
    ),
    557: (
        "System operation was verified through build checks, API health checks, database synchronization, seeded "
        "property retrieval, live ETH-price retrieval, and an on-chain escrow submission. The verified escrow "
        "test produced blockchain booking ID 1 and transaction hash "
        "0x0c2c342780cc1f6fe98bdac1b7de04dc2f82b3d0a8bad38dad217bce8e2a0866."
    ),
    559: "",
}

for idx, text in replacements.items():
    if idx < len(doc.paragraphs):
        set_paragraph_text(doc.paragraphs[idx], text)

# Insert objective-achievement subsection after the Chapter Five introduction.
intro = doc.paragraphs[527]
after = insert_paragraph_after(
    intro,
    "The following objective-achievement matrix links each research objective to the implemented result and "
    "measurable evidence obtained from the project.",
    style="Normal",
)
heading = insert_paragraph_after(after, "5.1 Achievement of Research Objectives", style="Heading 2")

rows = [
    [
        "Research objective",
        "How the objective was achieved",
        "Numerical evidence",
    ],
    [
        "1. Analyse vulnerabilities and financial friction points in Kigali residential booking.",
        "The study identified the main weaknesses of the traditional broker-led process and translated them "
        "into system requirements. The implemented design addresses deposit theft, double-booking, weak refund "
        "processes, dispute bias, identity fraud, owner impersonation, and suspicious rapid activity.",
        "7 risk/friction points mapped to system controls; 6 automated fraud-detection rules implemented.",
    ],
    [
        "2. Design and develop a Smart Contract-based escrow architecture.",
        "The RealEstateEscrow contract was implemented to lock deposits, require tenant and owner confirmation "
        "before release, allow timeout refund, support dispute resolution, and deduct a platform fee only after "
        "successful completion.",
        "393-line Solidity escrow contract; 2.5% platform fee; 1 verified on-chain escrow booking; transaction "
        "hash recorded.",
    ],
    [
        "3. Build a user-friendly decentralized booking application.",
        "The prototype provides a responsive PWA/web dApp for browsing properties, viewing details, booking, "
        "payments, wallet connection, KYC, messaging, notifications, and admin management. Non-crypto users are "
        "supported through MoMo and card payment flows.",
        "32 frontend pages; 12 reusable components; 5 supported languages; 3 payment methods; 6 seeded Kigali "
        "properties.",
    ],
    [
        "4. Evaluate effectiveness, security, and user acceptance in preventing deposit fraud.",
        "Technical evaluation was completed through compilation, build checks, API verification, database "
        "synchronization, blockchain escrow testing, and security-layer review. A user acceptance evaluation "
        "framework was prepared for future field testing.",
        "Backend errors reduced from about 60 to 0; 15 route files; 80+ API endpoints; 27 database models; "
        "UAT framework for 12-20 users.",
    ],
]

table = insert_table_after(heading, rows=len(rows), cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
widths = [2500, 4300, 2700]
for row_idx, row in enumerate(rows):
    for col_idx, value in enumerate(row):
        cell = table.cell(row_idx, col_idx)
        set_cell_text(
            cell,
            value,
            bold=row_idx == 0,
            color="FFFFFF" if row_idx == 0 else None,
            size=8 if row_idx > 0 else 9,
        )
        cell.width = widths[col_idx]
        if row_idx == 0:
            shade_cell(cell, "1F4E79")
        elif row_idx % 2 == 0:
            shade_cell(cell, "EAF2F8")

caption = insert_paragraph_after(
    heading,
    "Table 2: Achievement of Research Objectives with Implementation Evidence",
    style="Caption",
)
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
heading._p.addnext(table._tbl)
table._tbl.addnext(caption._p)

doc.save(DOCX_PATH)
print(f"Updated {DOCX_PATH}")
