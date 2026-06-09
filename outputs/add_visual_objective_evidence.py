from pathlib import Path
from textwrap import wrap

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


DOCX_PATH = Path(r"E:\realestate_booking\bachelor_of_technology_capstone_project.docx")
ASSET_DIR = Path(r"E:\realestate_booking\outputs\objective_evidence_assets")
ASSET_DIR.mkdir(parents=True, exist_ok=True)


def font(size, bold=False):
    candidates = [
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def draw_wrapped(draw, text, xy, width, fnt, fill, line_gap=6):
    x, y = xy
    avg_char = max(1, fnt.getlength("abcdefghijklmnopqrstuvwxyz") / 26)
    max_chars = max(18, int(width / avg_char))
    for line in wrap(text, max_chars):
        draw.text((x, y), line, font=fnt, fill=fill)
        y += fnt.size + line_gap
    return y


def create_objective4_figure(path):
    w, h = 1500, 900
    img = Image.new("RGB", (w, h), "#F7F4EA")
    d = ImageDraw.Draw(img)
    ink = "#10202B"
    teal = "#0D7C78"
    green = "#20A26B"
    gold = "#D8A135"
    red = "#B9574B"
    muted = "#5D716D"

    d.rectangle([0, 0, w, 80], fill="#09212A")
    d.text((48, 22), "Objective 4 Evidence: Evaluation of Effectiveness, Security, and User Acceptance", font=font(30, True), fill="white")
    d.text((48, 95), "The evaluation objective was addressed through technical verification, security controls, and a prepared user study framework.", font=font(24), fill=ink)

    cards = [
        ("Technical verification", green, "Frontend production build generated 32 pages and a PWA service worker. Backend API health, property retrieval, ETH price retrieval, and escrow submission were verified."),
        ("Security evaluation", teal, "JWT, RBAC, Zod validation, 2FA, KYC, file upload checks, rate limiting, fraud detection, ReentrancyGuard, and Pausable contract controls were implemented."),
        ("User acceptance instruments", gold, "EVALUATION_FRAMEWORK.md defines UAT methodology, SUS survey, perceived-security questionnaire, post-test interviews, and task-completion logging."),
        ("Comparative analysis", red, "The framework compares the smart-contract system against traditional agents on deposit safety, transparency, dispute handling, and trust."),
    ]
    x_positions = [70, 425, 780, 1135]
    for idx, (title, color, body) in enumerate(cards):
        x = x_positions[idx]
        y = 190
        d.rounded_rectangle([x, y, x + 295, y + 360], radius=18, fill="white", outline="#C9D8D2", width=2)
        d.rectangle([x, y, x + 295, y + 10], fill=color)
        d.text((x + 24, y + 38), title, font=font(24, True), fill=ink)
        draw_wrapped(d, body, (x + 24, y + 95), 245, font(20), muted, line_gap=8)

    metrics = [
        ("60 -> 0", "TypeScript errors reduced"),
        ("80+", "API endpoints"),
        ("27", "Database models"),
        ("12-20", "Planned UAT users"),
    ]
    for idx, (value, label) in enumerate(metrics):
        x = 155 + idx * 320
        y = 650
        d.rounded_rectangle([x, y, x + 250, y + 120], radius=14, fill="#09212A")
        d.text((x + 28, y + 22), value, font=font(34, True), fill="#20A26B")
        draw_wrapped(d, label, (x + 28, y + 70), 200, font(18), "white", line_gap=4)

    d.text((48, 835), "Source: Project implementation files, FINAL_REPORT.md, OBJECTIVES_ASSESSMENT.md, and EVALUATION_FRAMEWORK.md.", font=font(17), fill=muted)
    img.save(path)


def create_objective_evidence_flow(path):
    w, h = 1500, 760
    img = Image.new("RGB", (w, h), "#F3F7F5")
    d = ImageDraw.Draw(img)
    ink = "#10202B"
    teal = "#0D7C78"
    green = "#20A26B"
    gold = "#D8A135"
    red = "#B9574B"
    muted = "#5D716D"
    d.text((55, 45), "From Research Objectives to Implemented Evidence", font=font(34, True), fill=ink)
    d.text((55, 94), "Each objective produced a visible system artifact or verification output.", font=font(23), fill=muted)
    items = [
        ("Analyse risk", red, "7 booking risks mapped\n6 fraud rules"),
        ("Build escrow", teal, "393-line contract\n1 verified escrow tx"),
        ("Deliver dApp", green, "32 pages\n5 languages\n3 payment methods"),
        ("Evaluate", gold, "Build/API/DB checks\nUAT framework\nSUS + security survey"),
    ]
    for i, (title, color, body) in enumerate(items):
        x = 90 + i * 350
        y = 235
        d.rounded_rectangle([x, y, x + 270, y + 230], radius=18, fill="white", outline="#C9D8D2", width=2)
        d.rectangle([x, y, x + 270, y + 12], fill=color)
        d.text((x + 28, y + 45), title, font=font(27, True), fill=ink)
        draw_wrapped(d, body, (x + 28, y + 100), 210, font(23), muted, line_gap=9)
        if i < 3:
            d.line([x + 286, y + 115, x + 330, y + 115], fill=teal, width=5)
            d.polygon([(x + 330, y + 115), (x + 313, y + 104), (x + 313, y + 126)], fill=teal)
    d.text((55, 645), "This evidence moves the results chapter beyond claims: each objective is supported by code modules, screenshots, metrics, or evaluation instruments.", font=font(21, True), fill=ink)
    img.save(path)


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p = new_p
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def insert_picture_after(paragraph, image_path, width_inches=5.9):
    tmp = paragraph._parent.add_paragraph()
    tmp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = tmp.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))
    paragraph._p.addnext(tmp._p)
    return tmp


def set_body_style(paragraph):
    for run in paragraph.runs:
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(32, 32, 32)


objective4 = ASSET_DIR / "objective4_evaluation_evidence.png"
flow = ASSET_DIR / "objective_evidence_flow.png"
create_objective4_figure(objective4)
create_objective_evidence_flow(flow)

doc = Document(DOCX_PATH)
anchor = None
for para in doc.paragraphs:
    if "Achievement of Research Objectives with Implementation Evidence" in para.text:
        anchor = para
        para.text = "Table 5.1: Achievement of Research Objectives with Implementation Evidence"
        break
if anchor is None:
    raise RuntimeError("Objective evidence table caption not found")

current = insert_paragraph_after(anchor, "5.2 Visual Evidence Supporting Objective Achievement", style="Heading 2")

current = insert_paragraph_after(
    current,
    "The evidence for the objectives is not limited to the objective matrix. The implementation produced visible "
    "system artifacts, interface screens, data models, smart-contract verification outputs, and evaluation "
    "instruments. These artifacts demonstrate that the objectives moved from research intentions to implemented "
    "and testable project results.",
    style="Normal",
)
set_body_style(current)

current = insert_picture_after(current, flow, 6.3)
current = insert_paragraph_after(current, "Figure 5.1: Summary of objective evidence produced by the system implementation.", style="Caption")
current.alignment = WD_ALIGN_PARAGRAPH.CENTER

current = insert_paragraph_after(
    current,
    "For the first and second objectives, the strongest evidence is architectural and transactional. The complete "
    "database model shows that booking risks were translated into concrete data structures, while the smart-contract "
    "escrow implementation provides an auditable mechanism for locking, releasing, refunding, and disputing deposits.",
    style="Normal",
)
set_body_style(current)
current = insert_picture_after(current, Path(r"E:\realestate_booking\COMPLETE_ERD.png"), 5.8)
current = insert_paragraph_after(current, "Figure 5.2: Database and relationship evidence supporting vulnerability analysis and escrow workflow design.", style="Caption")
current.alignment = WD_ALIGN_PARAGRAPH.CENTER

current = insert_paragraph_after(
    current,
    "For the third objective, the user-facing evidence is shown through the implemented property-detail and wallet "
    "screens. These screens demonstrate that tenants can view property information, inspect location details, and "
    "continue toward payment or wallet-based escrow actions through a usable web interface.",
    style="Normal",
)
set_body_style(current)
current = insert_picture_after(current, Path(r"E:\realestate_booking\thesis_screenshots\04-property-detail.png"), 5.8)
current = insert_paragraph_after(current, "Figure 5.3: Property-detail interface evidence for the user-friendly dApp objective.", style="Caption")
current.alignment = WD_ALIGN_PARAGRAPH.CENTER

current = insert_picture_after(current, Path(r"E:\realestate_booking\thesis_screenshots\08-wallet.png"), 5.8)
current = insert_paragraph_after(current, "Figure 5.4: Wallet and payment interface evidence supporting decentralized booking access.", style="Caption")
current.alignment = WD_ALIGN_PARAGRAPH.CENTER

current = insert_paragraph_after(
    current,
    "For the fourth objective, the evaluation evidence combines technical checks, security controls, and planned "
    "user-acceptance instruments. The project verifies technical operation through build checks, API checks, database "
    "synchronization, and escrow transaction evidence. It also prepares structured evaluation instruments, including "
    "UAT scenarios, SUS usability measurement, perceived-security questions, and comparison against traditional "
    "intermediary booking.",
    style="Normal",
)
set_body_style(current)
current = insert_picture_after(current, objective4, 6.3)
current = insert_paragraph_after(current, "Figure 5.5: Evidence framework for Objective 4 evaluation of effectiveness, security, and user acceptance.", style="Caption")
current.alignment = WD_ALIGN_PARAGRAPH.CENTER

current = insert_paragraph_after(
    current,
    "The admin dashboard provides additional evidence for Objective 4 because it exposes the operational controls "
    "needed to monitor bookings, users, property approvals, disputes, fraud alerts, commissions, reports, and other "
    "platform activities. This supports the evaluation claim that the system is not only a tenant-facing prototype "
    "but also an administrable platform with mechanisms for security oversight and dispute governance.",
    style="Normal",
)
set_body_style(current)
current = insert_picture_after(current, Path(r"E:\realestate_booking\thesis_screenshots\06-admin.png"), 5.8)
current = insert_paragraph_after(current, "Figure 5.6: Admin dashboard evidence for monitoring, governance, and evaluation of platform activity.", style="Caption")
current.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.save(DOCX_PATH)
print(f"Added visual evidence figures to {DOCX_PATH}")
