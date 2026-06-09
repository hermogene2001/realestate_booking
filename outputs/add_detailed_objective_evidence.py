from pathlib import Path

from docx import Document
from docx.shared import Pt


DOCX_PATH = Path(r"E:\realestate_booking\bachelor_of_technology_capstone_project.docx")


def set_cell_text(cell, text, size=7.5, bold_first=False):
    cell.text = ""
    lines = text.split("\n")
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        run = paragraph.add_run(line)
        run.font.size = Pt(size)
        if bold_first and index == 0:
            run.bold = True


doc = Document(DOCX_PATH)

for table in doc.tables:
    header = " | ".join(cell.text.strip() for cell in table.rows[0].cells)
    if "Research objective" in header and "Numerical evidence" in header:
        set_cell_text(table.rows[0].cells[2], "Measured evidence and proof", size=9, bold_first=True)

        set_cell_text(
            table.rows[1].cells[2],
            "Numbers: 7 risk/friction points mapped; 6 fraud-detection rules implemented.\n"
            "Evidence: OBJECTIVES_ASSESSMENT.md links deposit theft, double-booking, refund weakness, dispute bias, identity fraud, owner impersonation, and rapid suspicious activity to system controls.\n"
            "System proof: booking.service.ts validates overlapping bookings; bookingTimeout.worker.ts handles timeout cancellation/refund logic; fraud.service.ts implements rule-based detection; kyc.routes.ts and propertyDocument.service.ts support identity and ownership verification.",
        )

        set_cell_text(
            table.rows[2].cells[2],
            "Numbers: 393-line RealEstateEscrow.sol contract; 2.5% platform fee; 1 verified escrow booking recorded on-chain.\n"
            "Evidence: blockchain/contracts/RealEstateEscrow.sol implements createBooking, confirmHandover, cancelBooking, raiseDispute, resolveDispute, refundTenant, and withdrawFees.\n"
            "Verification proof: contract address 0x5FbDB2315678afecb367f032d93F642f64180aa3; blockchain booking ID 1; transaction hash 0x0c2c342780cc1f6fe98bdac1b7de04dc2f82b3d0a8bad38dad217bce8e2a0866.",
        )

        set_cell_text(
            table.rows[3].cells[2],
            "Numbers: 32 frontend pages; 12 reusable components; 5 supported languages; 3 payment methods; 6 seeded Kigali properties.\n"
            "Evidence: frontend/src/app contains tenant, owner, and admin routes for property browsing, detail viewing, dashboard, bookings, payments, wallet, KYC, messages, notifications, and admin management.\n"
            "Usability proof: frontend/src/locales provides English, Kinyarwanda, French, Swahili, and Arabic translations; PropertyMap and LocationPicker support exact house location; payment flows support ETH, MTN MoMo, and card.",
        )

        set_cell_text(
            table.rows[4].cells[2],
            "Numbers: backend errors reduced from about 60 to 0 during implementation; 15 route files; 80+ API endpoints; 27 database models; UAT framework targets 12-20 users.\n"
            "Evidence: FINAL_REPORT.md records API health checks, frontend production build, database synchronization, property retrieval, ETH price retrieval, and escrow submission verification.\n"
            "Evaluation proof: EVALUATION_FRAMEWORK.md provides UAT scenarios, SUS usability survey, perceived-security questionnaire, comparative analysis matrix, technical metrics, and a 2-week data-collection schedule.",
        )
        break
else:
    raise RuntimeError("Objective achievement table not found")

doc.save(DOCX_PATH)
print("Detailed evidence added to each objective")
